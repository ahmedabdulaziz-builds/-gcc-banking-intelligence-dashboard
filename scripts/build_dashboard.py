"""Build the reusable GCC banking KPI comparison dashboard."""

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_FILE = PROJECT_ROOT / "data" / "processed" / "gcc_banking_verified_kpis.csv"
OUTPUT_FILE = PROJECT_ROOT / "outputs" / "excel" / "gcc_banking_dashboard.xlsx"

SHEET_NAMES = [
    "README",
    "Clean KPI Database",
    "Bank Comparison",
    "Profitability",
    "Balance Sheet Growth",
    "Asset Quality",
    "Capital & Liquidity",
    "Charts",
    "Analyst Notes",
    "Source Log",
]

CATEGORY_METRICS = {
    "Profitability": [
        "Net Profit",
        "Operating Income",
        "Net Interest Income",
        "Non-Interest Income",
        "Net Interest Margin",
        "RoTE",
        "Cost-to-Income Ratio",
    ],
    "Balance Sheet Growth": [
        "Total Assets",
        "Loans / Advances",
        "Gross Loans",
        "Customer Deposits",
        "Loan Growth",
        "Deposit Growth",
    ],
    "Asset Quality": ["NPL Ratio", "Coverage Ratio", "Cost of Risk"],
    "Capital & Liquidity": [
        "CET1 Ratio",
        "Capital Adequacy Ratio",
        "Liquidity Coverage Ratio",
        "Advances-to-Deposit Ratio",
    ],
}

CHART_METRICS = [
    ("Net Profit", "Net Profit by bank"),
    ("Operating Income", "Operating Income by bank"),
    ("Total Assets", "Total Assets by bank"),
    ("Customer Deposits", "Customer Deposits by bank"),
    ("Net Interest Margin", "Net Interest Margin by bank"),
    ("Cost-to-Income Ratio", "Cost-to-Income Ratio by bank"),
    ("NPL Ratio", "NPL Ratio by bank"),
    ("Coverage Ratio", "Coverage Ratio by bank"),
    ("CET1 Ratio", "CET1 Ratio by bank"),
    ("Liquidity Coverage Ratio", "LCR by bank"),
]

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_YELLOW = "FFF2CC"
WHITE = "FFFFFF"
TEXT_GREY = "666666"
THIN_GREY = Side(style="thin", color="B7B7B7")
WARNING = "Source-extracted figures pending final manual review. Do not treat this workbook as final verified analysis."


def read_source_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    """Read the required KPI CSV while preserving header and row order."""
    if not path.exists():
        raise FileNotFoundError(f"Required input file not found: {path}")
    with path.open(newline="", encoding="utf-8-sig") as input_file:
        reader = csv.DictReader(input_file)
        headers = reader.fieldnames or []
        rows = list(reader)
    if not headers:
        raise ValueError(f"Input file has no header row: {path}")
    return headers, rows


def numeric_value(value: str):
    """Convert a source value for derived workbook views without changing the source sheet."""
    try:
        return float(value.replace(",", "").strip())
    except (AttributeError, ValueError):
        return None


def style_title(cell) -> None:
    cell.font = Font(size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")


def style_header_row(sheet, row_number: int) -> None:
    for cell in sheet[row_number]:
        if cell.value is None:
            continue
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GREY)


def add_warning(sheet, row_number: int = 2, end_column: int = 6) -> None:
    sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=end_column)
    cell = sheet.cell(row=row_number, column=1, value=WARNING)
    cell.font = Font(bold=True, color="9C6500")
    cell.fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row_number].height = 32


def set_column_widths(sheet, maximum: int = 45) -> None:
    """Size columns from visible content while avoiding excessively wide sheets."""
    for column_cells in sheet.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[letter].width = min(max(width + 2, 12), maximum)


def add_table(sheet, name: str, header_row: int, last_row: int, last_column: int) -> None:
    """Add an Excel table to a populated range for filters and banded rows."""
    if last_row <= header_row:
        return
    ref = f"A{header_row}:{get_column_letter(last_column)}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def build_lookup(rows: list[dict[str, str]]) -> dict[tuple[str, str], dict[str, str]]:
    """Index source rows by bank and metric for comparison sheets."""
    return {(row["Bank"].strip(), row["Metric"].strip()): row for row in rows}


def ordered_banks(rows: list[dict[str, str]]) -> list[str]:
    """Preserve the source order of bank names."""
    return list(dict.fromkeys(row["Bank"].strip() for row in rows if row["Bank"].strip()))


def metric_unit(rows: list[dict[str, str]], metric: str) -> str:
    units = list(dict.fromkeys(row["Unit"].strip() for row in rows if row["Metric"].strip() == metric and row["Unit"].strip()))
    return units[0] if len(units) == 1 else " / ".join(units)


def format_value_cell(cell, unit: str) -> None:
    if unit in {"%", "% YTD"}:
        cell.number_format = '0.0"%"'
    elif unit == "bps":
        cell.number_format = '0.0" bps"'
    else:
        cell.number_format = "#,##0.0##"


# README section: explain purpose, workflow, source, and review limitations.
def populate_readme(sheet, rows: list[dict[str, str]]) -> None:
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "AI-Assisted GCC Banking Intelligence Dashboard"
    style_title(sheet["A1"])
    sheet.row_dimensions[1].height = 28

    banks = ", ".join(ordered_banks(rows)) or "No banks found"
    periods = ", ".join(dict.fromkeys(row["Period"] for row in rows if row["Period"])) or "No periods found"
    content = [
        ("Dashboard purpose", "Compare FAB and ENBD across profitability, balance sheet growth, asset quality, capital strength, liquidity, and efficiency."),
        ("Questions answered", "Which bank is larger, growing faster, more profitable, more efficient, better provisioned, and more strongly capitalised or liquid?"),
        ("Banks in source", banks),
        ("Periods in source", periods),
        ("Data source", INPUT_FILE.relative_to(PROJECT_ROOT).as_posix()),
        ("Workflow used", "Official source documents -> KPI extraction -> structured CSV -> Excel comparison views and charts -> final manual source review."),
        ("Review warning", WARNING),
        ("Use of figures", "Values in analytical tabs are derived directly from the source CSV. Missing metrics are left blank; no figures are invented."),
        ("AI disclosure", "AI-assisted workflow supported extraction, structuring, and drafting; final figures and conclusions require manual verification."),
    ]
    for row_number, (label, value) in enumerate(content, start=3):
        sheet.cell(row=row_number, column=1, value=label).font = Font(bold=True, color=NAVY)
        value_cell = sheet.cell(row=row_number, column=2, value=value)
        value_cell.alignment = Alignment(wrap_text=True, vertical="top")
        if label == "Review warning":
            value_cell.fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
            value_cell.font = Font(bold=True, color="9C6500")
    sheet.column_dimensions["A"].width = 23
    sheet.column_dimensions["B"].width = 105
    sheet.freeze_panes = "A3"


# Source database section: write every CSV field exactly as supplied.
def populate_clean_database(sheet, headers: list[str], rows: list[dict[str, str]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    style_header_row(sheet, 1)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    add_table(sheet, "CleanKPIDatabaseTable", 1, sheet.max_row, sheet.max_column)
    set_column_widths(sheet)


# Comparison section: create a pivot-style metric-by-bank view from source rows.
def populate_comparison(sheet, rows: list[dict[str, str]], banks: list[str], lookup) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(banks) + 2)
    sheet["A1"] = "Bank Comparison"
    style_title(sheet["A1"])
    add_warning(sheet, 2, len(banks) + 2)
    headers = ["Metric", "Unit", *banks]
    sheet.append(headers)
    header_row = 3

    metrics = list(dict.fromkeys(row["Metric"].strip() for row in rows if row["Metric"].strip()))
    for metric in metrics:
        unit = metric_unit(rows, metric)
        values = [f"{metric} ({unit})" if unit else metric, unit]
        for bank in banks:
            source_row = lookup.get((bank, metric))
            values.append(numeric_value(source_row["Value"]) if source_row else None)
        sheet.append(values)
        for cell in sheet[sheet.max_row][2:]:
            format_value_cell(cell, unit)

    style_header_row(sheet, header_row)
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    add_table(sheet, "BankComparisonTable", header_row, sheet.max_row, sheet.max_column)
    set_column_widths(sheet)


# Category sections: focus each finance topic on its requested KPI set.
def populate_category_sheet(sheet, title: str, metrics: list[str], rows, banks, lookup) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(banks) + 4)
    sheet["A1"] = title
    style_title(sheet["A1"])
    add_warning(sheet, 2, len(banks) + 4)
    headers = ["Metric", "Unit", *banks, "Period", "Review Status"]
    sheet.append(headers)
    header_row = 3

    for metric in metrics:
        unit = metric_unit(rows, metric)
        source_rows = [lookup.get((bank, metric)) for bank in banks]
        present_rows = [row for row in source_rows if row]
        values = [metric, unit]
        values.extend(numeric_value(row["Value"]) if row else None for row in source_rows)
        values.append(" / ".join(dict.fromkeys(row["Period"] for row in present_rows)))
        values.append(" / ".join(dict.fromkeys(row["Verification Status"] for row in present_rows)) or "Not found")
        sheet.append(values)
        for cell in sheet[sheet.max_row][2 : 2 + len(banks)]:
            format_value_cell(cell, unit)

    style_header_row(sheet, header_row)
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    table_name = title.replace("&", "And").replace(" ", "") + "Table"
    add_table(sheet, table_name, header_row, sheet.max_row, sheet.max_column)
    set_column_widths(sheet)


# Charts section: chart only metrics that exist in the source data.
def populate_charts(sheet, rows, banks, lookup) -> None:
    sheet.merge_cells("A1:P1")
    sheet["A1"] = "GCC Banking KPI Charts"
    style_title(sheet["A1"])
    add_warning(sheet, 2, 16)
    sheet["A3"] = "Charts use source-extracted values from the CSV. Review the Source Log and original documents before drawing final conclusions."
    sheet["A3"].font = Font(italic=True, color=TEXT_GREY)

    data_start = 90
    sheet.cell(data_start, 1, "Metric")
    for column, bank in enumerate(banks, start=2):
        sheet.cell(data_start, column, bank)
    style_header_row(sheet, data_start)

    positions = ["A5", "I5", "A21", "I21", "A37", "I37", "A53", "I53", "A69", "I69"]
    charts_added = 0
    for index, ((metric, title), position) in enumerate(zip(CHART_METRICS, positions), start=1):
        unit = metric_unit(rows, metric)
        data_row = data_start + index
        sheet.cell(data_row, 1, metric)
        has_data = False
        for column, bank in enumerate(banks, start=2):
            source_row = lookup.get((bank, metric))
            value = numeric_value(source_row["Value"]) if source_row else None
            sheet.cell(data_row, column, value)
            format_value_cell(sheet.cell(data_row, column), unit)
            has_data = has_data or value is not None
        if not has_data:
            continue

        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        data = Reference(sheet, min_col=2, max_col=1 + len(banks), min_row=data_row, max_row=data_row)
        categories = Reference(sheet, min_col=2, max_col=1 + len(banks), min_row=data_start, max_row=data_start)
        chart.add_data(data, from_rows=True, titles_from_data=False)
        chart.set_categories(categories)
        chart.title = title
        chart.y_axis.title = unit or "Value"
        chart.height = 7.5
        chart.width = 14
        chart.legend = None
        sheet.add_chart(chart, position)
        charts_added += 1

    if charts_added == 0:
        sheet["A5"] = "No chart-ready numeric values were found."
    sheet.sheet_state = "visible"
    sheet.row_dimensions[data_start].hidden = True
    for row_number in range(data_start + 1, data_start + 1 + len(CHART_METRICS)):
        sheet.row_dimensions[row_number].hidden = True
    sheet.freeze_panes = "A4"


# Analyst notes section: provide structured prompts without making conclusions.
def populate_analyst_notes(sheet) -> None:
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Analyst Notes Template"
    style_title(sheet["A1"])
    add_warning(sheet, 2, 6)
    sheet.append(["Section", "Observation / Question", "Fact / Calculation / Interpretation", "Source Reference", "Reviewer", "Status"])
    style_header_row(sheet, 3)
    sections = [
        "Profitability observations",
        "Balance sheet observations",
        "Asset quality observations",
        "Capital and liquidity observations",
        "Key risks to investigate",
        "Questions for interview discussion",
    ]
    for section in sections:
        sheet.append([section, "", "", "", "", "Open"])
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:F{sheet.max_row}"
    add_table(sheet, "AnalystNotesTable", 3, sheet.max_row, 6)
    set_column_widths(sheet)
    sheet.column_dimensions["B"].width = 48
    sheet.column_dimensions["D"].width = 35


# Source log section: retain the complete source trail required for each KPI.
def populate_source_log(sheet, rows: list[dict[str, str]]) -> None:
    headers = ["Bank", "Period", "Metric", "Value", "Unit", "Source Document", "Page/Section", "Verification Status", "Notes"]
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    style_header_row(sheet, 1)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    add_table(sheet, "SourceLogTable", 1, sheet.max_row, len(headers))
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    set_column_widths(sheet)


def build_dashboard() -> None:
    """Create, format, validate, and save the requested workbook."""
    headers, rows = read_source_csv(INPUT_FILE)
    banks = ordered_banks(rows)
    lookup = build_lookup(rows)

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {name: workbook.create_sheet(name) for name in SHEET_NAMES}

    populate_readme(sheets["README"], rows)
    populate_clean_database(sheets["Clean KPI Database"], headers, rows)
    populate_comparison(sheets["Bank Comparison"], rows, banks, lookup)
    for sheet_name, metrics in CATEGORY_METRICS.items():
        populate_category_sheet(sheets[sheet_name], sheet_name, metrics, rows, banks, lookup)
    populate_charts(sheets["Charts"], rows, banks, lookup)
    populate_analyst_notes(sheets["Analyst Notes"])
    populate_source_log(sheets["Source Log"], rows)

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_view.zoomScale = 90
    workbook["README"]["B9"].comment = Comment(
        "Final manual review should reconcile each KPI to the cited source page before recruiter or interview use.",
        "Finance source discipline",
    )

    workbook.save(OUTPUT_FILE)

    # Re-open once to catch malformed workbook structures before reporting success.
    checked = load_workbook(OUTPUT_FILE, data_only=False)
    if checked.sheetnames != SHEET_NAMES:
        raise RuntimeError("Workbook validation failed: unexpected sheet structure.")
    checked.close()
    print(f"Created dashboard workbook at {OUTPUT_FILE}")


if __name__ == "__main__":
    build_dashboard()
